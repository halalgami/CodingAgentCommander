#!/usr/bin/env python3
"""Raid heroine prompt sheet: docs/raid-art/heroine-prompts.xlsx + .csv.

Reads the real fal job files, so the sheet always matches what was generated.
Run from the repo root: python3 scripts/raidart/sheet_prompts.py
"""
import csv, io, json, os, shutil, sys
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(__file__))
import prep  # LOOKS, KEEP, LEGEND, GEAR, PROPS, IDS, POSES

CONCEPT = "docs/superpowers/concepts/2026-10-09-dashboard/fal"
ASSETS = "frontend/src/assets/raid"
OUT = "docs/raid-art"
HEROINES = [  # mirrors frontend/src/lib/raid/cast.js
    ("opus", "Archmage", "#e0a33a", "halo"),
    ("sonnet", "Spellblade", "#ef7a62", "halo"),
    ("haiku", "Ninja", "#4fb477", "wings"),
    ("glm", "Gunslinger", "#8f6ae0", "wings"),
    ("qwen", "Priestess", "#d45fb2", "halo"),
    ("kimi", "Mecha Pilot", "#36b5b0", "wings"),
]

stage1 = json.load(open(f"{CONCEPT}/stage1.json"))
stage2 = json.load(open(f"{CONCEPT}/stage2.json"))
x5 = json.load(open("scripts/raidart/heroines.json"))[0]["prompt"]
base = {j["out"].split("/")[-1][:-len("-stance.png")]: j["prompt"] for j in stage1 if "/raid/" in j["out"] and "-stance" in j["out"]}
pose = {p: next(j["prompt"] for j in stage2 if j["out"].endswith(f"opus-{p}.png")) for p in ("attack", "await", "hurt", "victory")}
look = lambda id, n: "recolour only: hair {}, eyes {}, outfit trim {}, and add {}".format(*prep.LOOKS[id][n])

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2B2F3A")
WRAP = Alignment(wrap_text=True, vertical="top")


def table(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


wb = Workbook()
ws = wb.active
ws.title = "Heroines"
hdr = ["Picture", "Used when the model name contains", "Heroine id", "Class", "Signature colour",
       "1. Base character prompt (text-to-image)", "2. Figure pass (edit, applied to every pose)",
       "Look 1 (recolour edit)", "Look 2 (recolour edit)", "Halo or wings (level 10)", "Familiar (level 3)"]
rows = [["", id, id, cls, tint, base[id], x5, look(id, 0), look(id, 1), aura, prep.GEAR[f"gear-familiar-{id}"]]
        for id, cls, tint, aura in HEROINES]
table(ws, hdr, rows, [16, 16, 10, 12, 10, 60, 50, 40, 40, 12, 36])
for i, (id, *_rest) in enumerate(HEROINES, 2):
    ws.row_dimensions[i].height = 120
    ws.cell(i, 5).fill = PatternFill("solid", fgColor=_rest[1][1:])
    im = Image.open(f"{ASSETS}/{id}-stance.webp").convert("RGBA")
    im.thumbnail((110, 155))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    pic = XLImage(buf)
    ws.add_image(pic, f"A{i}")
ws.append([])
ws.append(["", "Any other model", "(one of the six)", "Mercenary",
           "", "Borrows one of the six heroines, picked by a hash of the session id"])

ws = wb.create_sheet("Pose prompts")
table(ws, ["Pose", "When it shows", "Prompt (edit of the stance image)"], [
    ["stance", "working / idle", "(the base character prompt itself)"],
    ["attack", "running a tool", pose["attack"]],
    ["hurt", "error / failed tool", pose["hurt"]],
    ["victory", "Quest Clear (finished)", pose["victory"]],
    ["await", "waiting for your input or a permission", pose["await"]],
], [10, 30, 110])

ws = wb.create_sheet("Looks & legend")
table(ws, ["Step", "Prompt"], [
    ["Shared prefix for every look and legend edit", prep.KEEP],
    ["Legendary outfit (level 20) suffix", prep.LEGEND],
    ["Look suffix pattern", "recolour only: hair <hair>, eyes <eyes>, outfit trim <trim>, and add <accessory>"],
], [40, 120])

ws = wb.create_sheet("Gear")
table(ws, ["File", "Prompt", "Style suffix"],
      [[k + ".webp", v, prep.PROPS.lstrip(", ")] for k, v in prep.GEAR.items()], [26, 70, 70])

ws = wb.create_sheet("Files to replace")
files = []
for id, cls, *_ in HEROINES:
    for p in prep.POSES:
        files.append([id, cls, p, f"{id}-{p}.webp", f"{id}-v1-{p}.webp", f"{id}-v2-{p}.webp", f"{id}-legend-{p}.webp"])
    files.append([id, cls, "familiar", f"gear-familiar-{id}.webp", "", "", ""])
table(ws, ["Heroine", "Class", "Pose", "Main file (required)", "Look 1 (optional)", "Look 2 (optional)",
           "Legend (optional)"], files, [10, 12, 10, 24, 24, 24, 26])

ws = wb.create_sheet("Pipeline")
table(ws, ["Item", "Value"], [
    ["Text-to-image model", "fal-ai/nano-banana-2 (aspect 3:4)"],
    ["Edit model", "fal-ai/nano-banana-2/edit"],
    ["Background removal", "fal-ai/birefnet/v2"],
    ["Order", "1. base prompt → stance  2. pose edits from the stance  3. figure pass on every pose  "
              "4. look 1, look 2 and legend edits from each figure-pass pose"],
    ["Final size", "transparent webp, trimmed, 420 px tall (gear 300 px); whole raid pack ≤ 4 MB"],
    ["Folder", "frontend/src/assets/raid/"],
    ["Replace from the app", "Settings → Raid heroines: click or drop an image on a cell, or Import zip / Export starter kit"],
], [24, 120])

os.makedirs(OUT, exist_ok=True)
wb.save(f"{OUT}/heroine-prompts.xlsx")
with open(f"{OUT}/heroine-prompts.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(hdr[1:])
    for r in rows:
        w.writerow(r[1:])
shutil.copy(f"{OUT}/heroine-prompts.csv", "raidkit/heroine-prompts.csv")  # embedded in the starter kit
print(f"wrote {OUT}/heroine-prompts.xlsx and .csv")
